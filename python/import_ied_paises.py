"""Descarga el Excel de inversión extranjera directa por país del Banco Central y genera
public/data/ied_paises.json, que usa la sección de IED (globo y barra por país).

Uso: npm run import-ied-paises            (descarga el Excel)
     npm run import-ied-paises -- archivo.xlsx   (usa un archivo ya descargado)

Corre solo en la sincronización diaria de GitHub (no necesita credenciales: el archivo es público).

Se lee la hoja "País (Activo y Pasivo)", que trae cuatro bloques (millones de US$, anual):
  - flujos y stock de inversión directa ACTIVA: de Chile hacia el exterior ("outward").
  - flujos y stock de inversión directa PASIVA: del exterior hacia Chile ("inward").
Cada bloque tiene filas de continente (columna B), sus países (columna C), "No asignados" y
"TOTAL". Antes de escribir se verifica que países, continentes y total cuadren; si no, el formato
cambió y el script termina con error sin tocar el JSON (el sitio sigue con los datos anteriores).
"""

import io
import json
import os
import re
import sys
import unicodedata
from datetime import date

import requests
from openpyxl import load_workbook

SOURCE_URL = (
    "https://si3.bcentral.cl/estadisticas/Principal1/excel/SE/BDP/Excel/Anual/IED_pais_sector_region.xlsx"
)
OUTPUT_PATH = os.path.join("public", "data", "ied_paises.json")
TOLERANCE = 1  # millones de US$: las sumas del Excel pueden diferir por redondeo

# Título de cada bloque (normalizado) -> (sentido, medida).
BLOCKS = {
    "flujos de inversion directa activa por pais": ("outward", "flows"),
    "stock de inversion directa activa por pais": ("outward", "stock"),
    "flujos de inversion directa pasiva por pais": ("inward", "flows"),
    "stock de inversion directa pasiva por pais": ("inward", "stock"),
}


def plain(text):
    """Minúsculas, sin tildes y con espacios simples."""
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    return re.sub(r"\s+", " ", text).strip()


def normalize(text):
    """Como plain, pero además sin notas al pie como "(4)" o "(**)"."""
    return plain(re.sub(r"\([^)]*\)", " ", str(text or "")))


def clean_name(text):
    """Nombre para mostrar: sin notas al pie, p. ej. "Estados Unidos (4)" -> "Estados Unidos"."""
    return re.sub(r"\s+", " ", re.sub(r"\([^)]*\)", " ", str(text))).strip()


def as_number(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def find_sheet(workbook):
    for sheet in workbook.worksheets:
        name = plain(sheet.title)
        if name.startswith("pais") and "activo" in name and "pasivo" in name:
            return sheet
    raise ValueError("No se encontró la hoja 'País (Activo y Pasivo)'.")


def year_columns(row):
    """{índice de columna: año} si la fila es la de encabezado de años."""
    years = {i: int(v) for i, v in enumerate(row) if as_number(v) is not None and 1990 < v < 2100 and float(v).is_integer()}
    return years if len(years) >= 2 else None


def parse_block(rows, start):
    """Lee un bloque desde su título: devuelve (años, datos, fila siguiente al TOTAL)."""
    index = start + 1
    while index < len(rows) and not year_columns(rows[index]):
        index += 1
    if index >= len(rows):
        raise ValueError("Un bloque no tiene la fila de años.")
    columns = year_columns(rows[index])
    years = [columns[i] for i in sorted(columns)]

    def values(row):
        out = [as_number(row[i]) if i < len(row) else None for i in sorted(columns)]
        return [round(v, 3) if v is not None else None for v in out]

    block = {"continents": [], "unassigned": None, "total": None}
    for index in range(index + 1, len(rows)):
        row = rows[index]
        continent = row[1] if len(row) > 1 else None
        country = row[2] if len(row) > 2 else None
        if continent and normalize(continent) == "total":
            block["total"] = values(row)
            return years, block, index + 1
        if continent and normalize(continent).startswith("no asignado"):
            block["unassigned"] = values(row)
        elif continent and any(v is not None for v in values(row)):
            block["continents"].append({"name": clean_name(continent), "total": values(row), "countries": []})
        elif country and block["continents"]:
            block["continents"][-1]["countries"].append({"name": clean_name(country), "values": values(row)})
    raise ValueError("Un bloque no termina en una fila TOTAL.")


def check_block(block, where):
    """Cada continente es la suma de sus países, y el total la suma de continentes y no asignados."""
    def close(a, b):
        return a is None or b is None or abs(a - b) <= TOLERANCE

    for continent in block["continents"]:
        if not continent["countries"]:
            continue
        for i, value in enumerate(continent["total"]):
            parts = [c["values"][i] or 0 for c in continent["countries"]]
            if not close(value, sum(parts)):
                raise ValueError(f"{where}: {continent['name']} no es la suma de sus países (año {i}).")
    for i, value in enumerate(block["total"]):
        parts = [c["total"][i] or 0 for c in block["continents"]] + [(block["unassigned"] or [0] * len(block["total"]))[i] or 0]
        if not close(value, sum(parts)):
            raise ValueError(f"{where}: el TOTAL no es la suma de continentes y no asignados (año {i}).")


def parse_workbook(source):
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = [list(row) for row in find_sheet(workbook).iter_rows(values_only=True)]
    data = {"inward": {}, "outward": {}}
    years = None
    for index, row in enumerate(rows):
        title = normalize(row[1] if len(row) > 1 else "")
        if title in BLOCKS:
            direction, measure = BLOCKS[title]
            block_years, block, _ = parse_block(rows, index)
            if years and block_years != years:
                raise ValueError("Los bloques no tienen los mismos años.")
            years = block_years
            check_block(block, f"{direction}/{measure}")
            data[direction][measure] = block
    missing = [f"{d}/{m}" for d, m in BLOCKS.values() if m not in data[d]]
    if missing:
        raise ValueError(f"Faltan bloques en el Excel: {', '.join(missing)}")
    return years, data


def download():
    response = requests.get(SOURCE_URL, timeout=120, headers={"User-Agent": "Chileconomics (datos publicos)"})
    response.raise_for_status()
    return io.BytesIO(response.content)


def main(args):
    source = args[0] if args else download()
    years, data = parse_workbook(source)
    payload = {
        "source": SOURCE_URL,
        "updated": date.today().isoformat(),
        "unit": "Millones de dólares",
        "note": "País del inversionista inmediato (puede diferir del país del controlador final).",
        "years": years,
        **data,
    }
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, separators=(",", ":"))
    stock = data["inward"]["stock"]["total"]
    print(f"IED por país: {years[0]}-{years[-1]}, stock en Chile {years[-1]}: {stock[-1]:,.0f} millones de US$")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except (requests.RequestException, ValueError) as error:
        print(f"No se actualizó {OUTPUT_PATH}: {error}", file=sys.stderr)
        sys.exit(1)
