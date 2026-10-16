"""Descarga las proyecciones de crecimiento del PIB de Chile del Banco Central (IPoM) y del FMI
(World Economic Outlook) y genera public/data/proyecciones.json, que usa la sección de proyecciones.

Uso: npm run import-proyecciones

Corre solo en la sincronización diaria de GitHub (no necesita credenciales: las fuentes son públicas).
La tercera fuente, la Encuesta de Expectativas Económicas (EEE), viene de la API del Banco Central
y la descarga sync_bcch_data.py (claves eee_pib_<año>).

- IPoM: el Banco Central publica con cada IPoM un Excel ("Proyecciones IPoM 2010-20xx") con todas
  las proyecciones desde 2010. El enlace cambia con cada informe, así que se busca en la página del
  IPoM. La fila "PIB" trae rangos ("1,5 / 2,5", "4,25-5,25") o un número (estimación puntual).
- FMI: la API SDMX del FMI tiene la edición vigente del WEO y algunas anteriores
  (WEO_2025_OCT_VINTAGE...). Como las ediciones viejas pueden dejar de estar, las que ya están en el
  JSON se conservan: el historial crece con cada edición nueva.

Si una fuente falla se conserva su versión anterior; si fallan las dos, no se escribe nada y el
script termina con error (el sitio sigue con los datos anteriores).
"""

import io
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
from datetime import date, datetime
from urllib.parse import quote

import requests
from openpyxl import load_workbook

OUTPUT_PATH = os.path.join("public", "data", "proyecciones.json")
IPOM_PAGE = "https://www.bcentral.cl/web/banco-central/areas/politica-monetaria/informe-de-politica-monetaria"
BCCH_SITE = "https://www.bcentral.cl"
IMF_API = "https://api.imf.org/external/sdmx/2.1"
IMF_SERIES = "CHL.NGDP_RPCH.A"  # Chile, PIB a precios constantes, variación anual (%)
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; Chileconomics/1.0; +https://chileconomics.cl)"}
TIMEOUT = 120
DECIMALS = 2

MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
          "octubre", "noviembre", "diciembre"]
MONTH_CODES = {"JAN": 1, "APR": 4, "JUL": 7, "OCT": 10}
NUMBER = r"-?\d+(?:[.,]\d+)?"


def month_label(prefix, year, month):
    return f"{prefix} {MONTHS[month - 1]} {year}"


def is_year(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer()
            and 2000 <= value <= 2100)


def to_number(text):
    return float(text.replace(",", "."))


def parse_range(value):
    """Celda de la fila PIB -> [mínimo, máximo] (iguales si es un número) o None.

    Formatos del Excel: "4,5 - 5,5", "4,25-5,25", "-2,5 / -1,5", "-0,5 /0", "5,2" o un número.
    """
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = round(float(value), DECIMALS)
        return [number, number]
    text = str(value).strip()
    match = re.fullmatch(rf"({NUMBER})\s*(?:/|-)\s*({NUMBER})", text)
    if match:
        low, high = sorted([to_number(match.group(1)), to_number(match.group(2))])
        return [round(low, DECIMALS), round(high, DECIMALS)]
    if re.fullmatch(NUMBER, text):
        number = round(to_number(text), DECIMALS)
        return [number, number]
    return None


def find_ipom_link(html):
    """Enlace al Excel "Proyecciones IPoM ..." dentro de la página del IPoM (URL absoluta)."""
    for href in re.findall(r'href="([^"]+)"', html):
        if "proyecciones ipom" in href.lower() and ".xlsx" in href.lower():
            return BCCH_SITE + quote(href, safe="/-._~") if href.startswith("/") else href
    return None


def parse_ipom(rows):
    """Filas de la hoja del Excel -> vintages [{date, label, values: {año: [mín, máx]}}].

    Cada bloque de columnas es un año proyectado (fila de años) y cada columna un IPoM (fila de
    fechas). Se usa la fila cuya primera celda es "PIB".
    """
    years_row = dates_row = pib_row = None
    for index, row in enumerate(rows):
        if pib_row is None and str(row[0] or "").strip().lower() == "pib":
            pib_row = index
        if years_row is None and sum(is_year(cell) for cell in row) >= 2:
            years_row = index
        if dates_row is None and sum(isinstance(cell, (datetime, date)) for cell in row) >= 3:
            dates_row = index
    if None in (years_row, dates_row, pib_row):
        raise ValueError("No se encontraron las filas de años, fechas o PIB en el Excel del IPoM")

    def cell(row, column):
        return rows[row][column] if column < len(rows[row]) else None

    by_date = {}
    target = None
    for column, published in enumerate(rows[dates_row]):
        # El año va en la primera columna de cada bloque; las columnas siguientes son del mismo año.
        if is_year(cell(years_row, column)):
            target = int(cell(years_row, column))
        if not isinstance(published, (datetime, date)):
            continue
        value = parse_range(cell(pib_row, column))
        if target is None or value is None:
            continue
        key = f"{published.year}-{published.month:02d}-01"
        vintage = by_date.setdefault(key, {
            "date": key,
            "label": month_label("IPoM", published.year, published.month),
            "values": {},
        })
        vintage["values"][str(target)] = value
    return sorted(by_date.values(), key=lambda vintage: vintage["date"])


def load_ipom():
    page = requests.get(IPOM_PAGE, headers=HEADERS, timeout=TIMEOUT)
    page.raise_for_status()
    link = find_ipom_link(page.text)
    if not link:
        raise ValueError("No se encontró el Excel de proyecciones en la página del IPoM")
    response = requests.get(link, headers=HEADERS, timeout=TIMEOUT)
    response.raise_for_status()
    book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    sheet = book["Tabla"] if "Tabla" in book.sheetnames else book.worksheets[0]
    vintages = parse_ipom([list(row) for row in sheet.iter_rows(values_only=True)])
    if len(vintages) < 10:
        raise ValueError(f"El Excel del IPoM trae solo {len(vintages)} informes: ¿cambió el formato?")
    return {"source": link, "vintages": vintages}


def weo_flows(structure_xml):
    """Dataflows del WEO en el catálogo del FMI: [(id, versión)], sin repetir."""
    return sorted(set(re.findall(r'id="(WEO[A-Z0-9_]*)"\s+version="([^"]+)"', structure_xml)))


def parse_weo(data_xml):
    """Respuesta SDMX -> ({año: valor}, fecha de publicación 'YYYY-MM-DD' o None)."""
    root = ET.fromstring(data_xml)
    values = {}
    published = None
    for element in root.iter():
        tag = element.tag.split("}")[-1]
        if tag == "DataSet" and element.get("PUBLICATION_DATE"):
            published = element.get("PUBLICATION_DATE")[:10]
        if tag == "Obs" and element.get("TIME_PERIOD") and element.get("OBS_VALUE") not in (None, "", "NaN"):
            values[element.get("TIME_PERIOD")[:4]] = round(float(element.get("OBS_VALUE")), DECIMALS)
    return values, published


def weo_vintage(flow, values, published):
    """Edición del WEO a partir del nombre del dataflow (WEO_2025_OCT_VINTAGE) o de la fecha de
    publicación (dataflow WEO, la edición vigente). Solo se guardan los años desde el de la edición."""
    match = re.fullmatch(r"WEO_(\d{4})_([A-Z]{3})_VINTAGE", flow)
    if match and match.group(2) in MONTH_CODES:
        year, month = int(match.group(1)), MONTH_CODES[match.group(2)]
    elif flow == "WEO" and published:
        year, month = int(published[:4]), int(published[5:7])
    else:
        return None
    projected = {key: value for key, value in values.items() if int(key) >= year}
    if not projected:
        return None
    return {"date": f"{year}-{month:02d}-01", "label": month_label("WEO", year, month), "values": projected}


def load_weo(previous):
    catalog = requests.get(f"{IMF_API}/dataflow/IMF.RES/all/all", headers=HEADERS, timeout=TIMEOUT)
    catalog.raise_for_status()
    flows = weo_flows(catalog.text)
    latest = max((version for flow, version in flows if flow == "WEO"),
                 key=lambda v: [int(part) for part in v.split(".")], default=None)
    wanted = [(flow, version) for flow, version in flows if flow != "WEO"]
    if latest:
        wanted.append(("WEO", latest))

    vintages = {vintage["date"]: vintage for vintage in previous}
    found = 0
    for flow, version in wanted:
        response = requests.get(f"{IMF_API}/data/IMF.RES,{flow},{version}/{IMF_SERIES}", headers=HEADERS,
                                timeout=TIMEOUT)
        if not response.ok:
            print(f"   FMI {flow} {version}: HTTP {response.status_code}")
            continue
        vintage = weo_vintage(flow, *parse_weo(response.text))
        if vintage:
            vintages[vintage["date"]] = vintage
            found += 1
            print(f"   FMI {flow} {version}: {vintage['label']}")
    if not found:
        raise ValueError("La API del FMI no entregó ninguna edición del WEO")
    return {"source": f"{IMF_API}/data/IMF.RES,WEO/{IMF_SERIES}", "vintages": sorted(vintages.values(), key=lambda v: v["date"])}


def load_previous():
    if not os.path.exists(OUTPUT_PATH):
        return {}
    with open(OUTPUT_PATH, encoding="utf-8") as file:
        return json.load(file)


def write_atomic(payload):
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    temporary = f"{OUTPUT_PATH}.tmp"
    with open(temporary, "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=1)
        file.write("\n")
    os.replace(temporary, OUTPUT_PATH)


def main():
    previous = load_previous()
    payload = {
        "note": "Crecimiento anual del PIB de Chile proyectado (%). IPoM: rango [mín, máx] del Banco Central; "
                "WEO: proyección puntual del FMI.",
    }
    failures = 0
    for key, loader in (("ipom", load_ipom), ("weo", lambda: load_weo(previous.get("weo", {}).get("vintages", [])))):
        try:
            payload[key] = loader()
            print(f"{key}: {len(payload[key]['vintages'])} ediciones, la última {payload[key]['vintages'][-1]['label']}")
        except Exception as error:  # red, formato o API: se conserva lo anterior
            failures += 1
            print(f"{key}: ERROR {error}")
            if previous.get(key):
                payload[key] = previous[key]
                print(f"{key}: se conserva la versión anterior")
    if failures == 2:
        print("Fallaron las dos fuentes: no se escribe el archivo.")
        return 1
    write_atomic(payload)
    print(f"Datos guardados en {OUTPUT_PATH}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
