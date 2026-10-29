"""Pruebas del importador de IED por país con un Excel armado en la prueba (mismo formato que el real)."""

import io
import os
import sys

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import import_ied_paises  # noqa: E402

YEARS = [2024, 2025]


def block(sheet, title, continents, unassigned):
    """Escribe un bloque como en el Excel: título, años, continentes (col. B), países (col. C)."""
    sheet.append([None, title])
    sheet.append([None, "(Millones de dólares)"])
    sheet.append([None, None, None, *YEARS])
    total = [0, 0]
    for name, countries in continents:
        values = [sum(v[i] for _, v in countries) for i in range(2)]
        sheet.append([None, name, None, *values])
        for country, country_values in countries:
            sheet.append([None, None, country, *country_values])
        total = [total[i] + values[i] for i in range(2)]
    sheet.append([None, "No asignados", None, *unassigned])
    total = [total[i] + unassigned[i] for i in range(2)]
    sheet.append([None, "TOTAL", None, *total])
    sheet.append([None, None, "Unión Europea (7)", 1, 1])
    sheet.append([None, "(*) Incluye títulos de capital"])


def sector_block(sheet, title, sectors, unassigned, break_total=False):
    """Bloque de sectores como en el Excel: título, unidad, fila "Ítem | años", sectores y TOTAL."""
    sheet.append([None, title])
    sheet.append([None, "(Millones de dólares)"])
    sheet.append([])
    sheet.append([None, "Ítem", *YEARS])
    for name, values in sectors:
        sheet.append([None, name, *values])
    sheet.append([None, "No asignados", *unassigned])
    total = [sum(v[i] for _, v in sectors) + unassigned[i] for i in range(2)]
    sheet.append([None, "TOTAL", total[0] + (50 if break_total else 0), total[1]])
    sheet.append([None, "(**) Incluye patrimonio e instrumentos de deuda"])


def workbook(break_total=False, break_sectors=False):
    book = Workbook()
    book.active.title = "Portada"
    sheet = book.create_sheet("País (Activo y Pasivo)")
    sheet.append([None, "PRINCIPIO ACTIVO/PASIVO (1) (2) (3)"])
    america = [("Brasil", [10, 12]), ("Estados Unidos (4)", [5, 6])]
    block(sheet, "FLUJOS DE INVERSIÓN DIRECTA ACTIVA POR PAÍS (*)", [("América", america)], [1, 1])
    block(sheet, "STOCK DE INVERSIÓN DIRECTA ACTIVA POR PAÍS (**)", [("América", america)], [2, 2])
    inward = [
        ("América", [("Canadá", [40, 47]), ("Resto América (8)", [3, 3])]),
        ("Europa", [("España", [20, 20])]),
        ("África", []),
    ]
    block(sheet, "FLUJOS DE INVERSIÓN  DIRECTA PASIVA POR PAIS (*) (***)", inward[:2], [1, 2])
    block(sheet, "STOCK DE INVERSIÓN DIRECTA PASIVA POR PAÍS (**) (***)", inward, [30, 33])
    sectors = book.create_sheet("Sector (Activo y Pasivo)")
    sectors.append([None, "PRINCIPIO ACTIVO/PASIVO (1) (2) (3) "])
    sector_block(sectors, "FLUJOS DE INVERSIÓN DIRECTA ACTIVA POR SECTOR ECONÓMICO (*)", [("Minería", [1, 1])], [0, 0])
    sector_block(sectors, "STOCK DE INVERSIÓN DIRECTA ACTIVA POR SECTOR ECONÓMICO (**)",
                 [("Minería", [4, 5]), ("Servicios financieros", [10, 11])], [3, 4])
    sector_block(sectors, "STOCK DE INVERSIÓN DIRECTA PASIVA POR SECTOR ECONÓMICO (**)",
                 [("Minería", [60, 70]), ("Comercio", [3, 3])], [30, 30], break_total=break_sectors)
    if break_total:
        for row in sheet.iter_rows():
            if row[1].value == "TOTAL":
                row[3].value += 100
    data = io.BytesIO()
    book.save(data)
    data.seek(0)
    return data


def test_lee_los_cuatro_bloques():
    years, data = import_ied_paises.parse_workbook(workbook())
    assert years == YEARS
    stock = data["inward"]["stock"]
    assert [c["name"] for c in stock["continents"]] == ["América", "Europa", "África"]
    assert stock["continents"][0]["countries"][0] == {"name": "Canadá", "values": [40, 47]}
    assert stock["unassigned"] == [30, 33]
    assert stock["total"] == [93, 103]
    assert data["outward"]["stock"]["continents"][0]["countries"][1]["name"] == "Estados Unidos"


def test_continente_sin_paises_queda_con_su_total():
    _, data = import_ied_paises.parse_workbook(workbook())
    africa = data["inward"]["stock"]["continents"][2]
    assert africa["countries"] == []


def test_la_fila_union_europea_no_se_cuenta():
    _, data = import_ied_paises.parse_workbook(workbook())
    names = [c["name"] for cont in data["inward"]["stock"]["continents"] for c in cont["countries"]]
    assert "Unión Europea" not in names


def test_si_no_cuadra_el_total_falla():
    with pytest.raises(ValueError, match="TOTAL"):
        import_ied_paises.parse_workbook(workbook(break_total=True))


def test_normaliza_titulos_y_notas():
    assert import_ied_paises.normalize("STOCK DE INVERSIÓN  DIRECTA PASIVA POR PAÍS (**) (***)") == (
        "stock de inversion directa pasiva por pais"
    )
    assert import_ied_paises.clean_name("Estados Unidos (4)") == "Estados Unidos"


def test_lee_el_stock_por_sector_en_ambos_sentidos():
    _, data = import_ied_paises.parse_workbook(workbook())
    outward = data["outward"]["sectors"]
    assert outward["items"] == [{"name": "Minería", "values": [4, 5]}, {"name": "Servicios financieros", "values": [10, 11]}]
    assert outward["unassigned"] == [3, 4]
    assert outward["total"] == [17, 20]
    assert data["inward"]["sectors"]["total"] == [93, 103]


def test_si_los_sectores_no_suman_el_total_falla():
    with pytest.raises(ValueError, match="sectores"):
        import_ied_paises.parse_workbook(workbook(break_sectors=True))
